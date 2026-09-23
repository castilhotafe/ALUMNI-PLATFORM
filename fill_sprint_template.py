from docx import Document
from docx.shared import Pt
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from copy import deepcopy
from zipfile import ZipFile
from pathlib import Path
import hashlib
src=Path(r'C:\Users\konra\source\repos\ALUMNI-docs\ICT50220 INPRJP2 ASS1 P A1.docx')
out=Path('ICT50220 INPRJP2 ASS1 P A1 - Completed Sprint Plan.docx')
d=Document(src); t=d.tables[0]
plan=Document(r'C:\Users\konra\source\repos\ALUMNI-docs\Sprint Plan User Projects updated.docx')
Path('sprint-plan-qa').mkdir(exist_ok=True)
Path('sprint-plan-qa/artifact.md').write_text('Reference: '+str(src)+'\nSHA256: '+hashlib.sha256(src.read_bytes()).hexdigest()+'\nPreserve source sections, margins, merged cells, gray header fills, single borders, headers and footers. Fill team, leader, date, goals, tasks and contingencies. Clone task rows as needed for seven issues. Remove empty paragraph padding within filled cells; allow row growth. Renderer unavailable in this environment.\n',encoding='utf-8')
def put(c,text,bold=False):
    c.text=text
    c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in c.paragraphs:
        p.style=d.styles['Table Paragraph'] if 'Table Paragraph' in d.styles else d.styles['Normal']
        p.paragraph_format.left_indent=Pt(4)
        p.paragraph_format.right_indent=Pt(4)
        p.paragraph_format.space_before=Pt(3)
        p.paragraph_format.space_after=Pt(3)
        p.paragraph_format.line_spacing=1.0
        for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(11);r.bold=bold
put(t.cell(1,0),'Team Members:',True)
put(t.cell(1,1),'Jose Luis Quintero Ochoa\nEddrik Rei Uy Cana\nJuan David Garcia\nKonrad Krzyzkowiak\nMarcos de Castilho Junior')
put(t.cell(1,2),'Team Leader:',True);put(t.cell(1,3),'Konrad Krzyzkowiak')
put(t.cell(2,2),'Date:',True);put(t.cell(2,3),'17/09/2026\nSprint ends\n24/09/2026')
put(t.cell(3,0),'Team Goal(s):',True)
put(t.cell(3,1).merge(t.cell(3,3)),'USER PROJECT SHOWCASE & MANAGEMENT — EPIC 02\n'+next(p.text for p in plan.paragraphs if p.text.startswith('Let authenticated')))
# The template has five standard task rows and a merged separator row.
for _ in range(2):t.rows[10]._tr.addnext(deepcopy(t.rows[10]._tr))
for target,source in zip(list(t.rows)[6:13],list(plan.tables[0].rows)[1:]):
    for idx in range(3):put(target.cells[idx],source.cells[idx].text)
    target._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
put(t.rows[-1].cells[0],'Contingencies:',True)
put(t.rows[-1].cells[1], 'Account for unfinished work from the previous sprint, as discussed in the retrospective. Start database work first. Pages and test preparation can run alongside it once fields, routes, author permissions and visibility rules are agreed. Confirm authentication and layouts.platform are ready before integration.\nIf a task exceeds its estimate or blocks another member, Konrad will check progress and arrange support. On 22 September, review unfinished tasks and share or reassign work at risk. Keep 23 September for integration, reviews and fixes.')
# Use the template separator for the shared checkpoint without changing its structure.
put(t.rows[4].cells[0].merge(t.rows[4].cells[3]),'Proposed assignments and estimates — confirm at planning. By 21 September: at least one meaningful commit and 50% of agreed acceptance criteria per owner. By 24 September: all criteria met, tests passing, and work reviewed and merged.')
# Add notes in a cloned source row, preserving the template layout.
notes=deepcopy(t.rows[-1]._tr); t.rows[-1]._tr.addnext(notes)
put(t.rows[-1].cells[0],'Notes:',True)
put(t.rows[-1].cells[1],'Agenda: sprint goal and scope; tasks and responsibilities; workflow, dependencies and contingencies; questions and next steps.\nShare progress through WhatsApp or email on Monday and Wednesday, raise blockers promptly, and update GitHub issues. Konrad will review KPIs. Juan is asked to help with meeting minutes. Each PR needs two approvals before merging into dev. Run composer test and relevant frontend checks. Demonstrate the feature and discuss improvements on 24 September.\nEpic 02: github.com/castilhotafe/ALUMNI-PLATFORM/issues/22')
temp=Path('sprint-plan-qa/filled-working.docx');d.save(temp)
with ZipFile(src) as z,ZipFile(temp) as edited,ZipFile(out,'w') as final:
    for item in z.infolist():final.writestr(item,edited.read('word/document.xml') if item.filename=='word/document.xml' else z.read(item.filename))
with ZipFile(src) as z,ZipFile(out) as f:
    assert all(z.read(n)==f.read(n) for n in z.namelist() if n!='word/document.xml')
check=Document(out)
text=' '.join(c.text for row in check.tables[0].rows for c in row.cells)
assert all('#'+str(n) in text for n in range(24,31))
print(out.resolve())
print('Verified all seven tasks; all template package parts except document content preserved.')
